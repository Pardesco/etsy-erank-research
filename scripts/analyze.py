"""Turn eRank Keyword Tool exports into one research workbook + findings.json.

Usage (from the run folder, or pass it):
    python <skill>/scripts/analyze.py <run-folder> [--min-searches 100]

Reads
    02_seeds/seeds.csv                      seed,group,style,why
    02_seeds/relevance_terms.txt            optional, one word per line (empty = keep everything)
    02_seeds/exclude_terms.txt              optional, words that mean "different product"
    02_seeds/ip_terms_extra.txt             optional, extra brand/IP names found in review
    03_erank/exports/<slug>__keywords.csv   eRank Keyword Ideas export (CSV or XLSX)
    03_erank/exports/<slug>__listings.csv   eRank Top Listings export (CSV or XLSX), optional
    01_etsy/listings.json                   optional, phase-1 capture with styles
Writes
    04_results/research.xlsx, 04_results/findings.json
"""
import argparse
import csv
import json
import re
import statistics as st
import sys
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

warnings.simplefilter("ignore", FutureWarning)  # pandas concat notice on all-blank columns

HERE = Path(__file__).resolve().parent
DIGITAL = re.compile(r"\b(digital|printable|download|instant|downloadable|pdf|svg|png)\b", re.I)


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")


def terms(path):
    if not path.exists():
        return []
    return [t.strip().lower() for t in path.read_text(encoding="utf-8").splitlines() if t.strip() and not t.startswith("#")]


def word_regex(words):
    if not words:
        return None
    return re.compile(r"\b(?:" + "|".join(re.escape(w) for w in sorted(set(words), key=len, reverse=True)) + r")\b", re.I)


def num(v):
    """'$1,234.50' / '106%' / '< 20' / 'Unknown' -> float or None."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "").replace("%", "")
    m = re.match(r"^<\s*(\d+)", s)
    if m:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return None


def med(vals):
    vals = [v for v in vals if v is not None and not pd.isna(v)]
    return round(st.median(vals), 2) if vals else None


def per_1k(searches, comp):
    return round(searches * 1000 / comp, 2) if searches and comp else None


def read_table(path, header_words):
    """eRank exports start with a title row; find the real header row by its words."""
    if path.suffix.lower() in (".xlsx", ".xls"):
        raw = pd.read_excel(path, header=None)
    else:  # rows can differ in length (title row), so read with csv, not pandas
        with open(path, newline="", encoding="utf-8-sig") as fh:
            rows = [r for r in csv.reader(fh) if any(c.strip() for c in r)]
        width = max(len(r) for r in rows)
        raw = pd.DataFrame([r + [""] * (width - len(r)) for r in rows])
    for i in range(min(10, len(raw))):
        cells = [str(c).strip().lower() for c in raw.iloc[i].tolist()]
        if any(c in header_words for c in cells):
            df = raw.iloc[i + 1:].copy()
            df.columns = [str(c).strip() for c in raw.iloc[i].tolist()]
            return df.reset_index(drop=True)
    raise ValueError(f"{path.name}: no header row with any of {sorted(header_words)}")


def pick(columns, *rules):
    """First column whose lowercased name contains every word of a rule and none of its '!' words."""
    for rule in rules:
        want = [w for w in rule.split() if not w.startswith("!")]
        avoid = [w[1:] for w in rule.split() if w.startswith("!")]
        for c in columns:
            lc = c.lower()
            if all(w in lc for w in want) and not any(w in lc for w in avoid):
                return c
    return None


def load_keywords(path):
    df = read_table(path, {"keywords", "keyword"})
    cols = list(df.columns)
    m = {"Keyword": pick(cols, "keyword !difficulty !tag !count"),
         "Avg searches/mo": pick(cols, "search !google !trend"),
         "Avg clicks/mo": pick(cols, "click"),
         "CTR %": pick(cols, "ctr"),
         "KD (0-100)": pick(cols, "difficulty", "kd"),
         "Competition (listings)": pick(cols, "competition"),
         "Tag occurrences": pick(cols, "tag occ"),
         "Google searches/mo": pick(cols, "google")}
    if not m["Keyword"] or not m["Avg searches/mo"]:
        raise ValueError(f"{path.name}: missing keyword or searches column in {cols}")
    out = pd.DataFrame({k: (df[c] if c else None) for k, c in m.items()})
    out["Keyword"] = out["Keyword"].astype(str).str.strip().str.lower()
    out = out[out["Keyword"].ne("") & out["Keyword"].ne("nan")]
    for k in m:
        if k != "Keyword":
            out[k] = out[k].map(num)
    out.loc[out["Competition (listings)"] == 0, "Competition (listings)"] = None  # 0 = not measured
    return out


def load_listings(path):
    df = read_table(path, {"listing", "title", "listing title", "rank"})
    cols = list(df.columns)
    m = {"Rank": pick(cols, "rank"), "Title": pick(cols, "title", "listing !id !url"), "Shop": pick(cols, "shop"),
         "URL": pick(cols, "url", "link"), "Age (days)": pick(cols, "age"), "Views": pick(cols, "view !daily"),
         "Daily views": pick(cols, "daily"), "Est. sales": pick(cols, "sales"), "Price $": pick(cols, "price"),
         "Est. revenue $": pick(cols, "revenue"), "Hearts": pick(cols, "heart", "favor"), "Tags": pick(cols, "tag")}
    out = pd.DataFrame({k: (df[c] if c else None) for k, c in m.items()})
    for k in ("Rank", "Age (days)", "Views", "Daily views", "Est. sales", "Price $", "Est. revenue $", "Hearts"):
        out[k] = out[k].map(num)
    out["Title"] = out["Title"].fillna("").astype(str)
    out["Tags"] = out["Tags"].fillna("").astype(str)
    out = out[out["Title"].str.strip().ne("")]
    out["Format (inferred)"] = [("Digital" if DIGITAL.search(t + " " + g) else "Physical") for t, g in zip(out["Title"], out["Tags"])]
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("run", nargs="?", default=".")
    ap.add_argument("--min-searches", type=float, default=100)
    a = ap.parse_args()
    run = Path(a.run).resolve()
    seeds = pd.read_csv(run / "02_seeds" / "seeds.csv").fillna("")
    seeds["seed"] = seeds["seed"].str.strip().str.lower()
    exports = run / "03_erank" / "exports"
    rel_terms = terms(run / "02_seeds" / "relevance_terms.txt")
    relevant = word_regex(rel_terms)
    exclude = word_regex(terms(run / "02_seeds" / "exclude_terms.txt"))
    ip = word_regex(terms(HERE / "ip_terms.txt") + terms(run / "02_seeds" / "ip_terms_extra.txt"))

    kw_frames, lst_frames, seed_rows, missing = [], [], [], []
    for _, s in seeds.iterrows():
        sl = slug(s["seed"])
        kf = next(iter(sorted(exports.glob(f"{sl}__keywords.*"))), None)
        lf = next(iter(sorted(exports.glob(f"{sl}__listings.*"))), None)
        if not kf:
            missing.append(s["seed"])
            continue
        k = load_keywords(kf).assign(Seed=s["seed"], Group=s["group"])
        kw_frames.append(k)
        L = load_listings(lf).assign(Seed=s["seed"], Group=s["group"]) if lf else pd.DataFrame()
        if len(L):
            lst_frames.append(L)
        me = k[k["Keyword"] == s["seed"]]
        me = me.iloc[0] if len(me) else None
        dig = L[L["Format (inferred)"] == "Digital"] if len(L) else L
        phy = L[L["Format (inferred)"] == "Physical"] if len(L) else L
        seed_rows.append({
            "Group": s["group"], "Seed keyword": s["seed"], "Style": s.get("style", ""),
            "Avg searches/mo": me["Avg searches/mo"] if me is not None else None,
            "Avg clicks/mo": me["Avg clicks/mo"] if me is not None else None,
            "CTR %": me["CTR %"] if me is not None else None,
            "Competition (listings)": me["Competition (listings)"] if me is not None else None,
            "KD (0-100)": me["KD (0-100)"] if me is not None else None,
            "Searches per 1k listings": per_1k(me["Avg searches/mo"], me["Competition (listings)"]) if me is not None else None,
            "Related keywords": len(k),
            "Top listings": len(L),
            "Top listings: % digital (inferred)": round(100 * len(dig) / len(L)) if len(L) else None,
            "Median price digital $": med(dig["Price $"].tolist()) if len(L) else None,
            "Median price physical $": med(phy["Price $"].tolist()) if len(L) else None,
            "Median est. sales": med(L["Est. sales"].tolist()) if len(L) else None,
            "Total est. revenue $": round(L["Est. revenue $"].fillna(0).sum()) if len(L) else None,
        })
    if not kw_frames:
        sys.exit(f"No keyword exports found in {exports} for the seeds in seeds.csv")
    seeds_df = pd.DataFrame(seed_rows).sort_values("Searches per 1k listings", ascending=False, na_position="last")

    # ---------- keyword ideas, deduped across seeds ----------
    allk = pd.concat(kw_frames, ignore_index=True)
    agg = allk.groupby("Keyword").agg(**{
        "Avg searches/mo": ("Avg searches/mo", "max"), "Avg clicks/mo": ("Avg clicks/mo", "max"), "CTR %": ("CTR %", "max"),
        "Competition (listings)": ("Competition (listings)", "max"), "KD (0-100)": ("KD (0-100)", "max"),
        "Tag occurrences": ("Tag occurrences", "max"), "Google searches/mo": ("Google searches/mo", "max"),
        "Groups": ("Group", lambda g: ", ".join(sorted(set(g)))), "Seeds": ("Seed", lambda g: ", ".join(sorted(set(g)))),
        "Found under # seeds": ("Seed", "nunique")}).reset_index()
    agg["Searches per 1k listings"] = [per_1k(s, c) for s, c in zip(agg["Avg searches/mo"], agg["Competition (listings)"])]
    agg["On topic"] = agg["Keyword"].str.contains(relevant) if relevant else True
    agg["IP / licence risk"] = agg["Keyword"].str.contains(ip) if ip else False
    agg["Different product"] = agg["Keyword"].str.contains(exclude) if exclude else False
    agg["Words"] = agg["Keyword"].str.split().str.len()
    kw_df = agg.sort_values("Avg searches/mo", ascending=False)

    opp = kw_df[(kw_df["Avg searches/mo"] >= a.min_searches) & kw_df["Competition (listings)"].notna() & kw_df["On topic"]].copy()
    opp["Opportunity score"] = (opp["Searches per 1k listings"] * (1 - opp["KD (0-100)"].fillna(100) / 125)).round(2)
    opp = opp.sort_values("Opportunity score", ascending=False)
    clean = opp[~opp["IP / licence risk"] & ~opp["Different product"]]
    opp_cols = ["Keyword", "Avg searches/mo", "Competition (listings)", "KD (0-100)", "Searches per 1k listings",
                "Opportunity score", "Words", "Groups", "IP / licence risk", "Different product"]
    best_by_group = (clean.assign(Group=clean["Groups"].str.split(", ")).explode("Group")
                     .groupby("Group", group_keys=False).head(8)
                     [["Group"] + opp_cols[:7]].sort_values(["Group", "Opportunity score"], ascending=[True, False]))

    # Better phrasing: per seed, a related phrase with more demand and lower difficulty
    better = []
    for _, r in seeds_df.iterrows():
        s_seed, s_vol, s_kd = r["Seed keyword"], r["Avg searches/mo"], r["KD (0-100)"]
        if s_vol is None or pd.isna(s_vol):
            continue
        s_kd = 101 if s_kd is None or pd.isna(s_kd) else s_kd
        core = word_regex([w for w in s_seed.split() if w not in rel_terms and len(w) > 2])  # subject words, not "wall art"
        if not core:
            continue
        cand = clean[clean["Seeds"].str.split(", ").map(lambda ss: s_seed in ss) & clean["Keyword"].str.contains(core)
                     & (clean["Avg searches/mo"] > s_vol)
                     & (clean["KD (0-100)"].fillna(100) < s_kd)]
        if len(cand):
            c = cand.iloc[0]
            better.append({"Seed keyword": s_seed, "Seed searches/mo": s_vol, "Seed KD": r["KD (0-100)"], "Better phrase": c["Keyword"],
                           "Phrase searches/mo": c["Avg searches/mo"], "Phrase KD": c["KD (0-100)"],
                           "Phrase competition": c["Competition (listings)"], "Opportunity score": c["Opportunity score"]})
    better_df = pd.DataFrame(better)

    # ---------- top listings, tags, shops, prices ----------
    lst_df = pd.concat(lst_frames, ignore_index=True) if lst_frames else pd.DataFrame()
    tag_df = shop_df = price_df = pd.DataFrame()
    if len(lst_df):
        tag_count, tag_seeds = Counter(), defaultdict(set)
        for tags, seed in zip(lst_df["Tags"], lst_df["Seed"]):
            for t in {x.strip().lower() for x in re.split(r"[,;|]", tags) if x.strip()}:
                tag_count[t] += 1
                tag_seeds[t].add(seed)
        kwi = kw_df.set_index("Keyword")
        tag_df = pd.DataFrame([{"Tag": t, "Used by # top listings": n, "Across # seeds": len(tag_seeds[t]),
                                "eRank avg searches/mo": kwi["Avg searches/mo"].get(t), "eRank competition": kwi["Competition (listings)"].get(t)}
                               for t, n in tag_count.most_common() if n >= 3])
        if lst_df["Shop"].notna().any():
            shop_df = (lst_df.groupby("Shop").agg(**{"Top-listing appearances": ("Title", "count"),
                                                     "Seeds": ("Seed", lambda s: ", ".join(sorted(set(s)))),
                                                     "Median price $": ("Price $", "median"), "Est. sales (sum)": ("Est. sales", "sum"),
                                                     "Est. revenue $ (sum)": ("Est. revenue $", "sum")})
                       .reset_index().sort_values(["Top-listing appearances", "Est. revenue $ (sum)"], ascending=False))
        rows = []
        for (seed, group, fmt), sub in lst_df.groupby(["Seed", "Group", "Format (inferred)"]):
            p = sub["Price $"].dropna()
            if len(p):
                q = p.quantile([0.25, 0.5, 0.75]).round(2).tolist()
                rows.append({"Group": group, "Seed keyword": seed, "Format (inferred)": fmt, "Listings": len(sub),
                             "P25 $": q[0], "Median $": q[1], "P75 $": q[2], "Median est. sales": med(sub["Est. sales"].tolist())})
        price_df = pd.DataFrame(rows)

    # ---------- phase-1 styles vs eRank ----------
    style_df = pd.DataFrame()
    etsy_json = run / "01_etsy" / "listings.json"
    if etsy_json.exists():
        L1 = pd.DataFrame(json.loads(etsy_json.read_text(encoding="utf-8")).get("listings", []))
        if len(L1) and "style" in L1:
            L1["badged"] = L1["badges"].map(lambda b: bool(b))
            sd = seeds_df.set_index("Seed keyword")
            rows = []
            for style, sub in L1.groupby("style"):
                ss = [s for s in seeds[seeds["style"] == style]["seed"] if s in sd.index]
                best = max(ss, key=lambda s: sd.loc[s, "Avg searches/mo"] or 0) if ss else None
                rows.append({"Style (Etsy capture)": style, "Listings seen": len(sub), "Badged": int(sub["badged"].sum()),
                             "Bestseller": int(sub["badges"].map(lambda b: "Bestseller" in (b or [])).sum()),
                             "Popular now": int(sub["badges"].map(lambda b: "Popular now" in (b or [])).sum()),
                             "Median price $": med(sub["price"].tolist()), "eRank seeds": ", ".join(ss), "Best seed": best,
                             "Best seed searches/mo": sd.loc[best, "Avg searches/mo"] if best else None,
                             "Best seed searches per 1k listings": sd.loc[best, "Searches per 1k listings"] if best else None})
            style_df = pd.DataFrame(rows).sort_values("Badged", ascending=False)

    # ---------- findings.json ----------
    f = {
        "seeds_total": len(seeds), "seeds_with_exports": len(seed_rows), "seeds_missing_exports": missing,
        "keyword_rows": len(allk), "unique_keywords": len(kw_df), "listing_rows": len(lst_df),
        "opportunities": len(opp), "clean_opportunities": len(clean), "min_searches": a.min_searches,
        "ip_flagged_in_top50": int(opp.head(50)["IP / licence risk"].sum()),
        "seed_by_ratio": seeds_df.to_dict("records"),
        "top_opportunities": opp.head(60)[opp_cols].to_dict("records"),
        "top_clean": clean.head(40)[opp_cols[:8]].to_dict("records"),
        "best_by_group": best_by_group.to_dict("records"),
        "better_phrasing": better_df.to_dict("records"),
        "top_tags": tag_df.head(30).to_dict("records"),
        "top_shops": shop_df.head(15).to_dict("records"),
        "style_matrix": style_df.to_dict("records"),
    }
    if len(lst_df):
        f["digital_share_inferred"] = round(100 * (lst_df["Format (inferred)"] == "Digital").mean())
        f["median_price_digital"] = med(lst_df[lst_df["Format (inferred)"] == "Digital"]["Price $"].tolist())
        f["median_price_physical"] = med(lst_df[lst_df["Format (inferred)"] == "Physical"]["Price $"].tolist())
    out = run / "04_results"
    out.mkdir(exist_ok=True)
    (out / "findings.json").write_text(json.dumps(f, indent=1, default=lambda o: None if pd.isna(o) else str(o)), encoding="utf-8")

    # ---------- workbook ----------
    readme = [
        "Source: eRank Keyword Tool exports (Keyword Ideas + Top Listings), one pair per seed keyword.",
        f"{len(seed_rows)} of {len(seeds)} seeds have exports. {len(allk):,} keyword rows ({len(kw_df):,} unique), {len(lst_df):,} top-listing rows.",
        "Avg searches/clicks are eRank's monthly estimates. CTR over 100% means shoppers open more than one listing per search.",
        "Competition = Etsy listings eRank counts for the phrase. 0 or blank = not measured. KD = eRank keyword difficulty 0-100.",
        "Searches per 1k listings = demand vs supply. Opportunity score = that ratio x (1 - KD/125); missing KD counts as 100.",
        f"Opportunities: >= {a.min_searches:g} searches/mo, measured competition, on-topic (relevance_terms.txt).",
        "Clean Opportunities drops 'IP / licence risk' (brands, franchises, living artists) and 'Different product' (exclude_terms.txt). Flags are a keyword screen, not legal advice.",
        "Format (inferred) = Digital when the title or tags mention digital/printable/download; otherwise Physical.",
        "Est. sales / revenue are eRank estimates, not Etsy figures. Use them to compare, not to forecast.",
    ]
    sheets = [("README", pd.DataFrame({"About this workbook": readme})), ("Seeds", seeds_df)]
    if len(style_df):
        sheets.append(("Style Matrix", style_df))
    sheets += [("Clean Opportunities", clean[opp_cols[:8]]), ("Best by Group", best_by_group), ("Better Phrasing", better_df),
               ("All Opportunities (flagged)", opp[opp_cols]), ("Keyword Ideas (all)", kw_df)]
    if len(lst_df):
        sheets += [("Top Listings", lst_df), ("Top Tags", tag_df), ("Top Shops", shop_df), ("Price Bands", price_df)]
    xlsx = out / "research.xlsx"
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(xlsx)
    head = PatternFill("solid", fgColor="1F4E3D")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), head, Alignment(wrap_text=True, vertical="top")
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 140 if ws.title == "README" else 60)
        if ws.title != "README" and ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
    wb.save(xlsx)
    print(f"wrote {xlsx}")
    print(f"seeds {len(seed_rows)}/{len(seeds)} | keywords {len(kw_df):,} unique | listings {len(lst_df):,} | "
          f"opportunities {len(opp)} (clean {len(clean)}) | IP-flagged in top 50: {f['ip_flagged_in_top50']}")
    if missing:
        print("no exports yet for:", ", ".join(missing))


if __name__ == "__main__":
    main()
